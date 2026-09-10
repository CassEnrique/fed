import re
from copy import copy
from pathlib import Path

import openpyxl
from openpyxl import load_workbook
from openpyxl.formula.translate import Translator
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import column_index_from_string, get_column_letter


def normalize_rfc(value):
    return str(value).strip().upper() if value is not None else ""


def copy_row_style_only(ws, source_row, target_row, max_col):
    """
    Copia únicamente estilos de una fila a otra, sin copiar valores ni fórmulas.
    """
    # altura de fila
    if source_row in ws.row_dimensions:
        ws.row_dimensions[target_row].height = ws.row_dimensions[source_row].height

    for col in range(1, max_col + 1):
        src = ws.cell(row=source_row, column=col)
        dst = ws.cell(row=target_row, column=col)

        if src.has_style:
            dst._style = copy(src._style)

        dst.font = copy(src.font)
        dst.fill = copy(src.fill)
        dst.border = copy(src.border)
        dst.alignment = copy(src.alignment)
        dst.number_format = copy(src.number_format)
        dst.protection = copy(src.protection)

        # No copiar valor
        dst.value = None


def find_rfc_groups(ws, header_row=1, rfc_col=5):
    """
    Encuentra grupos consecutivos de RFC en columna E.
    Retorna lista de tuplas: (start_row, end_row, rfc)
    """
    data_start = header_row + 1
    max_row = ws.max_row

    if max_row < data_start:
        return []

    groups = []
    start = data_start
    current_rfc = normalize_rfc(ws.cell(row=data_start, column=rfc_col).value)

    for row in range(data_start + 1, max_row + 1):
        rfc = normalize_rfc(ws.cell(row=row, column=rfc_col).value)
        if rfc != current_rfc:
            groups.append((start, row - 1, current_rfc))
            start = row
            current_rfc = rfc

    groups.append((start, max_row, current_rfc))
    return groups


def insert_blank_styled_rows_per_rfc_group(
    file_path, sheet_name, header_row=1, rfc_col=5, rows_to_insert=3
):
    wb = load_workbook(file_path)
    ws = wb[sheet_name] if sheet_name in wb.sheetnames else wb.active

    max_col = ws.max_column
    groups = find_rfc_groups(ws, header_row=header_row, rfc_col=rfc_col)

    if not groups:
        wb.save(file_path)
        print("No se encontraron grupos.")
        return

    # Insertar de abajo hacia arriba
    for start_row, end_row, rfc in reversed(groups):
        insert_at = end_row + 1
        current_max_row = ws.max_row

        # Mover bloque inferior hacia abajo
        if insert_at <= current_max_row:
            ws.move_range(
                f"A{insert_at}:{ws.cell(row=current_max_row, column=max_col).coordinate}",
                rows=rows_to_insert,
                cols=0,
                translate=True,
            )

        # Usar la última fila del grupo como plantilla de estilo
        template_row = end_row

        # Crear filas nuevas vacías pero con estilo
        for i in range(rows_to_insert):
            new_row = insert_at + i
            copy_row_style_only(ws, template_row, new_row, max_col)

    wb.save(file_path)
    print(f"Archivo actualizado correctamente: {file_path}")


def limpiar_total_usd(
    archivo_xlsx,
    nombre_hoja=None,
    columna_moneda=16,  # P
    columna_total_usd=22,  # V
    fila_inicio=2,
):
    """
    Limpia el contenido de la columna Total USD cuando la columna Moneda es 'MXN'.

    - Modifica el mismo archivo.
    - No crea otro archivo.
    - Mantiene estilos de celda.
    """

    ruta = Path(archivo_xlsx)

    if not ruta.exists():
        raise FileNotFoundError(f"No existe el archivo: {ruta}")

    if ruta.suffix.lower() != ".xlsx":
        raise ValueError(f"El archivo debe ser .xlsx: {ruta}")

    wb = load_workbook(ruta)
    ws = wb[nombre_hoja] if nombre_hoja else wb.active

    ultima_fila = ws.max_row

    if ultima_fila < fila_inicio:
        wb.save(ruta)
        return 0

    celdas_limpiadas = 0

    for fila in range(ultima_fila, fila_inicio - 1, -1):
        valor_moneda = ws.cell(row=fila, column=columna_moneda).value

        if valor_moneda is None:
            continue

        moneda = str(valor_moneda).strip()

        if moneda == "MXN":
            celda = ws.cell(row=fila, column=columna_total_usd)

            if celda.value is not None:
                celda.value = None
                celdas_limpiadas += 1

    wb.save(ruta)
    return celdas_limpiadas


def actualizar_tipo_cambio(archivo_cambio, archivo_salida):
    # ============================================
    CSV_FILE = archivo_cambio
    EXCEL_FILE = archivo_salida
    HOJA = "Datos"

    COL_FECHA_CSV = "Fecha_Determinacion"
    COL_FECHA_EXCEL = "Fecha Registro"
    MAPEO_COLUMNAS = {"Tipo_de_Cambio": "Tipo de Cambio"}
    FORMATO_FECHA = "%Y-%m-%d"
    # ============================================

    print("Cargando archivos...")

    # 1. Leer CSV
    df_csv = pd.read_csv(CSV_FILE)
    df_csv[COL_FECHA_CSV] = pd.to_datetime(
        df_csv[COL_FECHA_CSV], format=FORMATO_FECHA, errors="coerce"
    )

    # 2. Crear diccionario de actualizaciones
    df_csv_unicos = df_csv.drop_duplicates(subset=[COL_FECHA_CSV], keep="last")
    df_csv_unicos["fecha_str"] = df_csv_unicos[COL_FECHA_CSV].dt.strftime("%Y-%m-%d")
    actualizaciones = df_csv_unicos.set_index("fecha_str")[
        list(MAPEO_COLUMNAS.keys())
    ].to_dict("index")

    print(f"CSV: {len(df_csv)} registros")
    print(f"Fechas disponibles para actualizar: {len(actualizaciones)}")

    # 3. Cargar Excel
    wb = load_workbook(EXCEL_FILE)
    ws = wb[HOJA] if HOJA in wb.sheetnames else wb.active

    # 4. Buscar columnas
    header_row = 1
    col_indices = {}
    for csv_col, excel_col in MAPEO_COLUMNAS.items():
        for cell in ws[header_row]:
            if cell.value == excel_col:
                col_indices[excel_col] = cell.column
                break

    fecha_col_idx = None
    for cell in ws[header_row]:
        if cell.value == COL_FECHA_EXCEL:
            fecha_col_idx = cell.column
            break

    if fecha_col_idx is None:
        print(f"ERROR: No se encontró la columna '{COL_FECHA_EXCEL}'")
        return

    # 5. Actualizar filas
    actualizados = 0
    for row in range(2, ws.max_row + 1):
        fecha_val = ws.cell(row=row, column=fecha_col_idx).value
        fecha_str = parse_fecha_excel(fecha_val)

        if fecha_str is None:
            continue

        if fecha_str in actualizaciones:
            for csv_col, excel_col in MAPEO_COLUMNAS.items():
                if excel_col in col_indices:
                    ws.cell(
                        row=row, column=col_indices[excel_col]
                    ).value = actualizaciones[fecha_str][csv_col]
            actualizados += 1

    print(f"Registros actualizados: {actualizados}")

    wb.save(EXCEL_FILE)
    print(f"Archivo guardado: {EXCEL_FILE}")


COLUMNAS_SUMAR = {
    "I": 9,
    "K": 11,
    "L": 12,
    "M": 13,
    "N": 14,
    "O": 15,
    "R": 18,
    "T": 20,
    "V": 22,
}

COL_POLIZA = 3
COL_RFC = 5
FILA_INICIO = 2


def copiar_estilo_base(ws, fila_origen, fila_destino, columnas):
    for col in columnas:
        origen = ws.cell(row=fila_origen, column=col)
        destino = ws.cell(row=fila_destino, column=col)

        if origen.has_style:
            destino._style = copy(origen._style)

        destino.font = copy(origen.font)
        destino.fill = copy(origen.fill)
        destino.border = copy(origen.border)
        destino.alignment = copy(origen.alignment)
        destino.number_format = copy(origen.number_format)
        destino.protection = copy(origen.protection)

    if fila_origen in ws.row_dimensions:
        ws.row_dimensions[fila_destino].height = ws.row_dimensions[fila_origen].height


def insertar_totales_por_rfc(archivo_xlsx, nombre_hoja=None):
    wb = load_workbook(archivo_xlsx)
    ws = wb[nombre_hoja] if nombre_hoja else wb.active

    ultima_fila_datos = ws.max_row

    fila_actual = FILA_INICIO
    inicio_grupo = FILA_INICIO
    rango_sumatorias = []

    columnas_a_formatear = list(COLUMNAS_SUMAR.values())

    while fila_actual <= ultima_fila_datos:
        rfc_actual = ws.cell(row=fila_actual, column=COL_RFC).value
        rfc_actual = str(rfc_actual).strip() if rfc_actual is not None else ""

        siguiente_fila = fila_actual + 1
        rfc_siguiente = ""

        if siguiente_fila <= ultima_fila_datos:
            rfc_siguiente = ws.cell(row=siguiente_fila, column=COL_RFC).value
            rfc_siguiente = (
                str(rfc_siguiente).strip() if rfc_siguiente is not None else ""
            )

        if (
            (rfc_actual != rfc_siguiente) or (siguiente_fila > ultima_fila_datos)
        ) and rfc_actual != "":
            fila_insertar = fila_actual + 1
            current_max_row = ws.max_row
            max_col = ws.max_column

            # ws.insert_rows(fila_insertar, amount=1)
            ws.move_range(
                f"A{fila_insertar}:{ws.cell(row=current_max_row, column=max_col).coordinate}",
                rows=1,
                cols=0,
                translate=True,
            )
            ultima_fila_datos += 1
            rango_sumatorias.append(fila_insertar)

            copiar_estilo_base(ws, fila_actual, fila_insertar, columnas_a_formatear)

            num_inicio_excel = inicio_grupo
            num_fin_excel = fila_actual

            for letra_col, idx_col in COLUMNAS_SUMAR.items():
                celda_formula = ws.cell(row=fila_insertar, column=idx_col)

                if idx_col > 15:
                    formula = f'=IF($P{num_fin_excel}="MXN", "", SUM({letra_col}{num_inicio_excel}:{letra_col}{num_fin_excel}))'
                else:
                    formula = f"=SUM({letra_col}{num_inicio_excel}:{letra_col}{num_fin_excel})"

                celda_formula.value = formula
                fuente = copy(celda_formula.font)
                fuente.bold = True
                celda_formula.font = fuente

            inicio_grupo = fila_insertar + 1
            fila_actual = fila_insertar + 1
        else:
            fila_actual += 1

    fila_final_totales = ultima_fila_datos + 2

    for letra_col, idx_col in COLUMNAS_SUMAR.items():
        if rango_sumatorias:
            refs = ",".join(f"{letra_col}{row}" for row in rango_sumatorias)
            ws.cell(row=fila_final_totales, column=idx_col).value = f"=SUM({refs})"

    borde_superior = Side(style="thin", color="000000")
    col_inicio = columnas_a_formatear[0]
    col_fin = columnas_a_formatear[-1]

    for col in range(col_inicio, col_fin + 1):
        celda = ws.cell(row=fila_final_totales, column=col)
        borde_actual = copy(celda.border)
        celda.border = Border(
            left=borde_actual.left,
            right=borde_actual.right,
            top=borde_superior,
            bottom=borde_actual.bottom,
            diagonal=borde_actual.diagonal,
            diagonal_direction=borde_actual.diagonal_direction,
            outline=borde_actual.outline,
            vertical=borde_actual.vertical,
            horizontal=borde_actual.horizontal,
        )

    wb.save(archivo_xlsx)


def escribir_formula_condicional(
    ws, fila_ini, fila_fin, col_destino, tiene_mxn, tiene_usd
):
    """
    Genera e inserta la fórmula de suma óptima en la columna destino.
    fila_ini y fila_fin están en base 1 de Excel/openpyxl.
    """
    celda_destino = ws.cell(row=fila_fin, column=col_destino)

    partes_formula = []

    # Si en el grupo hubo registros en MXN, sumamos la columna O
    if tiene_mxn:
        partes_formula.append(f"SUM(O{fila_ini}:O{fila_fin})")

    # Si en el grupo hubo registros en USD, sumamos la columna V
    if tiene_usd:
        partes_formula.append(f"SUM(V{fila_ini}:V{fila_fin})")

    # Construir fórmula final
    if len(partes_formula) == 2:
        formula = f"={partes_formula[0]}+{partes_formula[1]}"
    elif len(partes_formula) == 1:
        formula = f"={partes_formula[0]}"
    else:
        # Caso de respaldo
        formula = f"=SUM(M{fila_ini}:O{fila_fin})+SUM(T{fila_ini}:V{fila_fin})"

    # Solo cambia el contenido de la celda; el estilo se conserva
    celda_destino.value = formula


def sumar_por_grupos_multimoneda(archivo_xlsx, nombre_hoja=None):
    """
    Transpilación de la macro UNO a openpyxl.
    Modifica el mismo archivo .xlsx sin crear otro.
    """

    wb = load_workbook(archivo_xlsx)
    ws = wb[nombre_hoja] if nombre_hoja else wb.active

    # Configuraciones de columnas en base 1 para openpyxl
    COL_NUMERO = 3  # C
    COL_RFC = 5  # E
    COL_MONEDA = 16  # P
    COL_MONTO = 25  # Y

    # Fila donde inician los datos reales
    fila_inicio_datos = 2  # Fila 2 real en Excel

    last_row = ws.max_row

    grupo_inicio = fila_inicio_datos

    prev_numero = None
    prev_rfc = None

    tiene_mxn = False
    tiene_usd = False

    # Recorremos una fila extra para forzar el cierre del último bloque
    for fila in range(fila_inicio_datos, last_row + 2):
        if fila <= last_row:
            celda_num = ws.cell(row=fila, column=COL_NUMERO).value
            celda_rfc = ws.cell(row=fila, column=COL_RFC).value
            celda_moneda = ws.cell(row=fila, column=COL_MONEDA).value

            celda_num = str(celda_num).strip() if celda_num is not None else ""
            celda_rfc = str(celda_rfc).strip() if celda_rfc is not None else ""
            celda_moneda = (
                str(celda_moneda).strip().upper() if celda_moneda is not None else ""
            )

            # Si la fila está vacía (separadores), cerramos grupo si aplica
            if not celda_num and not celda_rfc:
                if prev_numero is not None or prev_rfc is not None:
                    escribir_formula_condicional(
                        ws, grupo_inicio, fila - 1, COL_MONTO, tiene_mxn, tiene_usd
                    )
                    prev_numero, prev_rfc = None, None
                    tiene_mxn, tiene_usd = False, False
                continue
        else:
            # Forzar cierre del último grupo
            celda_num, celda_rfc, celda_moneda = "", "", ""

        # Inicializar primer grupo válido
        if prev_numero is None and prev_rfc is None and (celda_num or celda_rfc):
            grupo_inicio = fila
            prev_numero = celda_num
            prev_rfc = celda_rfc
            tiene_mxn = celda_moneda == "MXN"
            tiene_usd = celda_moneda == "USD"
            continue

        # Detectar cambio de grupo o fin de datos
        if (celda_num != prev_numero or celda_rfc != prev_rfc) and (
            prev_numero is not None
        ):
            escribir_formula_condicional(
                ws, grupo_inicio, fila - 1, COL_MONTO, tiene_mxn, tiene_usd
            )

            # Resetear para nuevo grupo
            if fila <= last_row and (celda_num or celda_rfc):
                grupo_inicio = fila
                prev_numero = celda_num
                prev_rfc = celda_rfc
                tiene_mxn = celda_moneda == "MXN"
                tiene_usd = celda_moneda == "USD"
            else:
                prev_numero, prev_rfc = None, None
                tiene_mxn, tiene_usd = False, False
        else:
            # Seguimos en el mismo grupo
            if celda_moneda == "MXN":
                tiene_mxn = True
            elif celda_moneda == "USD":
                tiene_usd = True

    wb.save(archivo_xlsx)


TIPOS_VALIDOS = {"Egresos", "Ingresos", "Diario"}


def normalizar_numero_poliza(valor):
    if valor is None:
        return ""

    if isinstance(valor, bool):
        return str(valor).strip()

    if isinstance(valor, int):
        return str(valor)

    if isinstance(valor, float):
        return str(int(valor)) if valor.is_integer() else str(valor).strip()

    return str(valor).strip()


def asignar_consecutivos_polizas(archivo_xlsx, nombre_hoja=None):
    ruta = Path(archivo_xlsx)

    if not ruta.exists():
        raise FileNotFoundError(f"No existe el archivo: {ruta}")

    if ruta.suffix.lower() != ".xlsx":
        raise ValueError(f"El archivo debe ser .xlsx: {ruta}")

    wb = load_workbook(ruta)
    ws = wb[nombre_hoja] if nombre_hoja else wb.active

    COL_TIPO = 2
    COL_NUMERO = 3
    COL_RFC = 5
    COL_DESTINO = 26  # Z

    last_row = ws.max_row
    filas_datos = []

    for r in range(2, last_row + 1):
        tipo = ws.cell(row=r, column=COL_TIPO).value
        numero = ws.cell(row=r, column=COL_NUMERO).value
        rfc = ws.cell(row=r, column=COL_RFC).value

        tipo = str(tipo).strip() if tipo is not None else ""
        numero = normalizar_numero_poliza(numero)
        rfc = str(rfc).strip() if rfc is not None else ""

        if tipo in TIPOS_VALIDOS and numero != "":
            filas_datos.append(
                {"row_index": r, "tipo": tipo, "numero": numero, "rfc": rfc}
            )

    if not filas_datos:
        wb.save(ruta)
        return 0

    p_contador = 0
    diario_numero_mapeo = {}
    idx = 0
    total_filas = len(filas_datos)
    asignaciones = 0

    while idx < total_filas:
        actual = filas_datos[idx]
        tipo_actual = actual["tipo"]
        num_actual = actual["numero"]
        rfc_actual = actual["rfc"]

        if tipo_actual in {"Egresos", "Ingresos"}:
            end_idx = idx
            while (
                end_idx + 1 < total_filas
                and filas_datos[end_idx + 1]["tipo"] == tipo_actual
                and filas_datos[end_idx + 1]["rfc"] == rfc_actual
                and filas_datos[end_idx + 1]["numero"] == num_actual
            ):
                end_idx += 1

            p_contador += 1
            fila_destino = filas_datos[end_idx]["row_index"]

            celda_ref_edo = ws.cell(row=fila_destino, column=COL_DESTINO)
            celda_ref_edo.value = f"P.{p_contador}"
            celda_ref_edo.font = Font(
                bold=True, color="DC0042", size=12
            )  # Negritas y color rojo
            celda_ref_edo.alignment = Alignment(horizontal="center", vertical="center")

            asignaciones += 1
            idx = end_idx + 1

        else:  # Diario
            if num_actual in diario_numero_mapeo:
                poliza_asignada = diario_numero_mapeo[num_actual]
            else:
                p_contador += 1
                poliza_asignada = f"P.{p_contador}"
                diario_numero_mapeo[num_actual] = poliza_asignada

            end_idx = idx
            while (
                end_idx + 1 < total_filas
                and filas_datos[end_idx + 1]["tipo"] == "Diario"
                and filas_datos[end_idx + 1]["numero"] == num_actual
                and filas_datos[end_idx + 1]["rfc"] == rfc_actual
            ):
                end_idx += 1

            fila_destino = filas_datos[end_idx]["row_index"]
            celda_ref_edo = ws.cell(row=fila_destino, column=COL_DESTINO)
            celda_ref_edo.value = poliza_asignada
            celda_ref_edo.font = Font(
                bold=True, color="DC0042", size=12
            )  # Negritas y color rojo
            celda_ref_edo.alignment = Alignment(horizontal="center", vertical="center")
            asignaciones += 1
            idx = end_idx + 1

    wb.save(ruta)
    # return asignaciones


def aplicar_borde_inferior_folios(archivo_xlsx, nombre_hoja=None):
    """
    Aplica borde inferior a filas según cambio de folio o cambio de RFC especial.

    - Modifica el mismo archivo .xlsx
    - No crea otro archivo
    - Conserva fórmulas y estilos existentes, añadiendo solo borde inferior
    """

    wb = load_workbook(archivo_xlsx)
    ws = wb[nombre_hoja] if nombre_hoja else wb.active

    # Columnas en base 1 para openpyxl
    col_inicio = 1  # A
    col_fin = 26  # Z  (fiel al valor 25 del código UNO)
    col_folio = 3  # C
    col_rfc = 5  # E

    skip_rfc = {"XAXX010101000"}

    ultima_fila = ws.max_row

    if ultima_fila < 2:
        wb.save(archivo_xlsx)
        return

    # Borde inferior similar al original
    borde_inferior = Side(style="medium", color="000000")

    # Recorrer desde la fila 2 hasta la penúltima
    for fila in range(2, ultima_fila):
        rfc_actual = ws.cell(row=fila, column=col_rfc).value
        rfc_siguiente = ws.cell(row=fila + 1, column=col_rfc).value

        folio_actual = ws.cell(row=fila, column=col_folio).value
        folio_siguiente = ws.cell(row=fila + 1, column=col_folio).value

        rfc_actual = str(rfc_actual).strip() if rfc_actual is not None else ""
        rfc_siguiente = str(rfc_siguiente).strip() if rfc_siguiente is not None else ""

        folio_actual = str(folio_actual).strip() if folio_actual is not None else ""
        folio_siguiente = (
            str(folio_siguiente).strip() if folio_siguiente is not None else ""
        )

        poner_borde = False

        if rfc_actual in skip_rfc:
            # Lógica especial: si cambia el RFC en la siguiente fila
            if rfc_actual != rfc_siguiente:
                poner_borde = True
        else:
            # Lógica normal: si cambia el folio
            if folio_actual and folio_actual != folio_siguiente:
                poner_borde = True

        if poner_borde:
            for col in range(col_inicio, col_fin + 1):
                celda = ws.cell(row=fila, column=col)
                borde_actual = copy(celda.border)
                celda.border = Border(
                    left=borde_actual.left,
                    right=borde_actual.right,
                    top=borde_actual.top,
                    bottom=borde_inferior,
                    diagonal=borde_actual.diagonal,
                    diagonal_direction=borde_actual.diagonal_direction,
                    outline=borde_actual.outline,
                    vertical=borde_actual.vertical,
                    horizontal=borde_actual.horizontal,
                )

    # Caso especial: la última fila con folio lleva borde
    folio_final = ws.cell(row=ultima_fila, column=col_folio).value
    folio_final = str(folio_final).strip() if folio_final is not None else ""

    if folio_final:
        for col in range(col_inicio, col_fin + 1):
            celda = ws.cell(row=ultima_fila, column=col)
            borde_actual = copy(celda.border)
            celda.border = Border(
                left=borde_actual.left,
                right=borde_actual.right,
                top=borde_actual.top,
                bottom=borde_inferior,
                diagonal=borde_actual.diagonal,
                diagonal_direction=borde_actual.diagonal_direction,
                outline=borde_actual.outline,
                vertical=borde_actual.vertical,
                horizontal=borde_actual.horizontal,
            )

    wb.save(archivo_xlsx)


def main_iva_format_process(file_name):
    archivo_cambio = "cambio_obligaciones.csv"

    # Paso 1, limpiar la columna Total USD
    total = limpiar_total_usd(file_name)
    print(f"Se limpiaron {total} celdas en Total USD.")

    # Paso 2,
    actualizar_tipo_cambio(archivo_cambio, file_name)

    # Paso 3, separar por RFC
    insert_blank_styled_rows_per_rfc_group(
        file_path=file_name, sheet_name=None, header_row=1, rfc_col=5, rows_to_insert=3
    )

    # Paso 4,
    insertar_totales_por_rfc(file_name)

    # Paso 5,
    sumar_por_grupos_multimoneda(file_name)

    # Paso 6,
    _ = asignar_consecutivos_polizas(file_name)

    # Paso 7,
    aplicar_borde_inferior_folios(file_name)
