import datetime
import io
import os
import shutil
import tempfile

import openpyxl
from openpyxl.drawing.image import Image
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from PyQt5.QtCore import QFile, QIODevice
from PyQt5.QtGui import QPixmap


def obtener_imagen_desde_recursos(ruta_recurso):
    """
    Extrae imagen desde recursos Qt compilados (funciona en exe también)
    """
    temp_file = tempfile.NamedTemporaryFile(suffix=".png", delete=False)
    temp_path = temp_file.name
    temp_file.close()

    try:
        qfile = QFile(ruta_recurso)
        if not qfile.open(QIODevice.ReadOnly):
            raise FileNotFoundError(f"No se pudo abrir: {ruta_recurso}")

        byte_array = qfile.readAll()
        qfile.close()

        with open(temp_path, "wb") as f:
            f.write(bytes(byte_array))

        return temp_path

    except Exception as e:
        print(f"❌ Error al extraer imagen: {e}")
        raise


def actualizar_excel_con_imagen(ruta_archivo, suffix, ruta_imagen):
    try:
        # --- PASO 1: OBTENER VALOR DE LA FECHA (MODO LECTURA) ---wb = openpyxl.load_workbook(ruta_archivo, keep_vba=True, data_only=False)
        wb_lectura = openpyxl.load_workbook(ruta_archivo, data_only=True)
        sheet_lectura = wb_lectura.active

        nombre_archivo_mayus = os.path.basename(ruta_archivo).upper()

        # Lógica de prefijos y celdas
        if "V0" in nombre_archivo_mayus:
            prefijo = "INTEGRACION DE VENTAS COBRADAS A LA TASA 0% EXPORTACION"
            celda_fecha = "A8"
        elif "V16" in nombre_archivo_mayus:
            prefijo = "INTEGRACION DE VENTAS COBRADAS A LA TASA 16%"
            celda_fecha = "Q8"
        elif "IVA" in nombre_archivo_mayus:
            prefijo = "CEDULA DE VENTAS GRAVADAS AL 16%"
            celda_fecha = "A8"
        else:
            prefijo = "INTEGRACION DE VENTAS"
            celda_fecha = "A8"

        valor_fecha = sheet_lectura[celda_fecha].value
        wb_lectura.close()

        # --- PASO 2: PROCESAR EL TÍTULO ---
        mes_año_texto = ""
        if valor_fecha:
            meses_map = {
                "Jan": "ENERO",
                "Feb": "FEBRERO",
                "Mar": "MARZO",
                "Apr": "ABRIL",
                "May": "MAYO",
                "Jun": "JUNIO",
                "Jul": "JULIO",
                "Aug": "AGOSTO",
                "Sep": "SEPTIEMBRE",
                "Oct": "OCTUBRE",
                "Nov": "NOVIEMBRE",
                "Dec": "DICIEMBRE",
                "Ene": "ENERO",
                "Abr": "ABRIL",
                "Ago": "AGOSTO",
                "Dic": "DICIEMBRE",
            }
            if isinstance(valor_fecha, (datetime.datetime, datetime.date)):
                mes_ing = valor_fecha.strftime("%b")
                mes_año_texto = (
                    f"{meses_map.get(mes_ing, mes_ing).upper()} {valor_fecha.year}"
                )
            else:
                str_fecha = str(valor_fecha).strip()
                partes = (
                    str_fecha.split("/") if "/" in str_fecha else str_fecha.split("-")
                )
                if len(partes) >= 3:
                    mes_nombre = meses_map.get(
                        partes[1].strip(), partes[1].strip()
                    ).upper()
                    mes_año_texto = f"{mes_nombre} {partes[2].strip()[:4]}"

        if not mes_año_texto:
            mes_año_texto = "MAYO 2025"
        titulo_completo = f"{prefijo} {mes_año_texto}"

        # --- PASO 3: EDICIÓN REAL (MODO DISEÑO) ---
        # data_only=False para NO borrar fórmulas. keep_vba=True para no romper macros.
        wb = openpyxl.load_workbook(ruta_archivo, data_only=False, keep_vba=True)
        sheet = wb.active

        # ESTRATEGIA: En lugar de insert_rows, usamos move_range.
        # Esto mueve todo el contenido (fórmulas, estilos, etc.) hacia abajo.
        # 'translate=True' ajusta las fórmulas para que sigan apuntando a sus datos.
        max_r = sheet.max_row
        max_c = sheet.max_column
        rango_a_mover = f"A1:{get_column_letter(max_c)}{max_r}"
        sheet.move_range(rango_a_mover, rows=6, cols=0, translate=True)

        # 3.2 Insertar la imagen en el espacio vacío creado
        img = Image(ruta_imagen)
        img.anchor = "A1"
        sheet.add_image(img)

        # 3.3 Calcular posición del texto basándose en el ancho de la imagen
        ancho_px = img.width
        col_actual = 0
        acumulado_px = 0
        while acumulado_px < ancho_px:
            col_actual += 1
            col_l = get_column_letter(col_actual)
            ancho_col = sheet.column_dimensions[col_l].width or 9.14
            acumulado_px += ancho_col * 7.5

        letra_col_texto = get_column_letter(col_actual + 2)

        # 3.4 Escribir los encabezados en las nuevas filas vacías de arriba
        sheet[
            f"{letra_col_texto}2"
        ].value = "RONGTAI INDUSTRIAL DEVELOPMENT LEON S DE RL DE CV"
        sheet[f"{letra_col_texto}2"].font = Font(bold=True, size=12)

        sheet[f"{letra_col_texto}3"].value = titulo_completo
        sheet[f"{letra_col_texto}3"].font = Font(bold=False, size=11)

        # --- PASO 4: GUARDAR ---
        # wb.save(ruta_archivo)
        temp_path = ruta_archivo + ".tmp"
        wb.save(temp_path)
        wb.close()

        shutil.move(temp_path, ruta_archivo)
        print(f"Éxito: Se actualizó {nombre_archivo_mayus} correctamente.")
        print(f"✅ {os.path.basename(ruta_archivo)} actualizado correctamente.")

    except Exception as e:
        print(f"❌ Error al procesar {os.path.basename(ruta_archivo)}: {e}")
        print(f"Error crítico: {e}")
        import traceback

        traceback.print_exc()


def main_header_process(file_path, suffix, logo_recurso):
    """
    Args:
        file_path: Ruta del archivo Excel
        suffix: Sufijo
        logo_recurso: Ruta del recurso (ej: ':/icons/logo.png')
    """

    if os.path.exists(file_path):
        # actualizar_excel_con_imagen(file_path, suffix, logo)

        # Extraer la imagen de los recursos
        ruta_imagen_temporal = obtener_imagen_desde_recursos(logo_recurso)
        try:
            actualizar_excel_con_imagen(file_path, suffix, ruta_imagen_temporal)
        finally:
            # Limpiar el archivo temporal
            if os.path.exists(ruta_imagen_temporal):
                os.remove(ruta_imagen_temporal)

    else:
        print(f"No existe: {file_path}")
