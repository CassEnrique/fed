from copy import copy
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

# ====================== FUNCIONES ======================


def get_header(ws, header_row=1):
    """Obtiene cabeceras con su posición (nombre → columna)"""
    headers = {}
    for col in range(1, ws.max_column + 1):
        cell = ws.cell(row=header_row, column=col)
        val = str(cell.value).strip() if cell.value is not None else ""
        if val:
            headers[val] = col
    return headers


def compare_and_map_headers(ws_aax, ws_abx, header_row=1):
    """Compara cabeceras por nombre y crea un mapping {nombre: (col_aax, col_abx)}"""
    h_aax = get_header(ws_aax, header_row)
    h_abx = get_header(ws_abx, header_row)

    common = {}
    missing_in_aax = []
    missing_in_abx = []

    for name, col_abx in h_abx.items():
        if name in h_aax:
            common[name] = (h_aax[name], col_abx)
        else:
            missing_in_aax.append(name)

    for name in h_aax:
        if name not in h_abx:
            missing_in_abx.append(name)

    return common, missing_in_aax, missing_in_abx, h_aax, h_abx


def copy_cell_style(source_cell, target_cell):
    """Copia estilo de una celda a otra"""
    if source_cell.has_style:
        target_cell.font = copy(source_cell.font)
        target_cell.fill = copy(source_cell.fill)
        target_cell.border = copy(source_cell.border)
        target_cell.alignment = copy(source_cell.alignment)
        target_cell.number_format = source_cell.number_format


# ====================== EJECUCIÓN ======================


def ejecutar_proceso_unir_archivo(file_aax, file_abx, sheet_name, output_file):
    print("Cargando archivos...")
    wb_aax = openpyxl.load_workbook(file_aax)
    wb_abx = openpyxl.load_workbook(file_abx)

    if sheet_name:
        ws_aax = wb_aax[sheet_name]
        ws_abx = wb_abx[sheet_name]
        print(sheet_name)
    else:
        ws_aax = wb_aax.active
        ws_abx = wb_abx.active
        print(sheet_name)

    common, missing_in_aax, missing_in_abx, h_aax, h_abx = compare_and_map_headers(
        ws_aax, ws_abx
    )

    print("\n=== MAPPING DE COLUMNAS COMUNES ===")
    for name, (col_aax, col_abx) in common.items():
        print(
            f"  {name}: aax({get_column_letter(col_aax)}) ← abx({get_column_letter(col_abx)})"
        )

    print(f"\n✅ Columnas comunes: {len(common)}")
    if missing_in_aax:
        print(f"❌ Columnas en abx que NO existen en aax: {missing_in_aax}")
    if missing_in_abx:
        print(f"❌ Columnas en aax que NO existen en abx: {missing_in_abx}")

    if not common:
        print("❌ No hay columnas comunes. Abortando.")
        wb_aax.close()
        wb_abx.close()
        exit()

    # ====================== CREAR ARCHIVO MERGEADO ======================

    print("\nCreando archivo merged...")
    wb_merged = openpyxl.Workbook()
    ws_merged = wb_merged.active
    ws_merged.title = "Merged_Data"

    # 1. Copiar cabecera y todo el contenido de aax (con estilos)
    for row_idx in range(1, ws_aax.max_row + 1):
        for col_idx in range(1, ws_aax.max_column + 1):
            old_cell = ws_aax.cell(row=row_idx, column=col_idx)
            new_cell = ws_merged.cell(row=row_idx, column=col_idx, value=old_cell.value)
            copy_cell_style(old_cell, new_cell)

    # Copiar merged cells de aax
    for merged_range in ws_aax.merged_cells.ranges:
        ws_merged.merge_cells(str(merged_range))

    # Copiar anchos de columna y alturas de fila de aax
    for col_letter, dim in ws_aax.column_dimensions.items():
        ws_merged.column_dimensions[col_letter].width = dim.width
    for row_num, dim in ws_aax.row_dimensions.items():
        ws_merged.row_dimensions[row_num].height = dim.height

    # 2. Añadir registros de abx (solo columnas comunes)
    start_row = ws_aax.max_row + 1
    print(f"Añadiendo datos de abx desde la fila {start_row}...")

    for row_idx_abx in range(2, ws_abx.max_row + 1):  # Saltamos cabecera
        new_row = start_row + (row_idx_abx - 2)
        for name, (col_aax, col_abx) in common.items():
            old_cell = ws_abx.cell(row=row_idx_abx, column=col_abx)
            new_cell = ws_merged.cell(row=new_row, column=col_aax, value=old_cell.value)
            copy_cell_style(old_cell, new_cell)

    # Copiar merged cells de abx (ajustando filas)
    for merged_range in ws_abx.merged_cells.ranges:
        new_start = merged_range.min_row + start_row - 2
        new_end = merged_range.max_row + start_row - 2
        new_range = f"{get_column_letter(merged_range.min_col)}{new_start}:{get_column_letter(merged_range.max_col)}{new_end}"
        ws_merged.merge_cells(new_range)

    # Guardar
    print(f"Guardando archivo como {output_file}...")
    wb_merged.save(output_file)

    print(f"\n✅ Merge completado exitosamente!")
    print(f"   - Archivo original (aax): {file_aax}")
    print(f"   - Archivo añadido (abx): {file_abx}")
    print(f"   - Archivo resultado: {output_file}")
    print(f"   - Total filas en resultado: {ws_merged.max_row}")

    wb_aax.close()
    wb_abx.close()
    wb_merged.close()


def main_merge_process(self, first_path, second_path):
    # ====================== CONFIGURACIÓN ======================
    # path = directorio = os.path.dirname(first_path)
    path = str(Path(first_path).parent)

    first_path = Path(first_path).resolve()
    second_path = Path(second_path).resolve()

    sheet_name = None

    output_file = f"{path}/merge.xlsx"

    # ============================================================
    # VALIDACIÓN PREVIA (Aquí detona y sale antes de intentar nada)
    # Si no existen los insumos, lanzamos el error de inmediato
    # ============================================================
    if not first_path.exists() or not second_path.exists():
        faltante = first_path if not first_path.exists() else second_path
        msg = f"No se encontró el archivo necesario: {faltante}"
        self.text_console_log(msg, "ERROR")
        # Este RAISE detiene la función AQUÍ MISMO y va directo al except de la GUI
        raise FileNotFoundError(msg)

    # ============================================================
    # 3. Llamar a las siguientes funciones (ahora es seguro)
    # ============================================================
    try:
        ejecutar_proceso_unir_archivo(first_path, second_path, sheet_name, output_file)
        print("✅ Proceso Unir Archivos completado exitosamente")
        self.text_console_log(
            "✅ Proceso Unir Archivos completado exitosamente", "PROCESS"
        )
        self.message_box(
            "Unir Archivos", "Proceso Unir Archivos completado exitosamente...", "ok"
        )
    except FileNotFoundError as e:
        print(f"❌ Archivo no encontrado: {e}")
        self.text_console_log(f"❌ Archivo no encontrado: {e}", "INFO")
    except Exception as e:
        print(f"❌ Error inesperado: {e}")
        self.text_console_log(f"❌ Error inesperado: {e}", "ERROR")
