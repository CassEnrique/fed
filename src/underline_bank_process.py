import re
from datetime import datetime

import cv2
import fitz  # PyMuPDF
import numpy as np
from openpyxl import load_workbook

# import easyocr


def crear_reader():
    import easyocr

    return easyocr.Reader(["es", "en"], gpu=False)


def extraer_caracteres_alfabeticos(texto):
    patron = re.compile(r"[A-Za-z]+")
    coincidencias = patron.findall(texto)
    resultado = "".join(coincidencias)
    return resultado


def formatear_fecha(fecha_val, type_bank):
    if isinstance(fecha_val, datetime):
        return fecha_val.strftime("%d")  # Solo el día para buscar en el PDF
    if type_bank == "HSBC":
        return str(fecha_val).split("/")[0].zfill(2)
    if isinstance(fecha_val, str):
        try:
            objeto_fecha = datetime.strptime(fecha_val, "%d/%b/%Y")
            return objeto_fecha.strftime("%d")
        except ValueError:
            return str(fecha_val).split("/")[0].zfill(2)
    return str(fecha_val).split("/")[0].zfill(2)


def obtener_monto_por_poliza(archivo_base, currency, version, nombre_hoja, type_bank):
    rfc_excluidos = ["XAXX010101000", "HMI950125KG8"]
    wb_iva = load_workbook(archivo_base, data_only=True)
    ws_iva = wb_iva[nombre_hoja] if nombre_hoja in wb_iva.sheetnames else wb_iva.active
    datos_monto = []
    compensation = 0

    for row in ws_iva.iter_rows(min_row=2, values_only=True):
        # ... (Asignaciones de columnas se mantienen igual)
        bank_val = "HSBC" if currency == "MXN" else "BASE"
        fecha_val = row[0] if len(row) > 0 else None
        rfc_val = row[4] if len(row) > 4 else None
        moneda_val = row[15 - compensation] if len(row) > 15 else None
        monto_val = row[24 - compensation] if len(row) > 24 else None
        edo_cta_val = row[25 - compensation] if len(row) > 25 else None

        if version == "v0":
            monto_val = row[14 - compensation] if len(row) > 14 else None
            edo_cta_val = row[15 - compensation] if len(row) > 15 else None
            bank_val = row[16] if len(row) > 16 else None

        if version == "v16":
            col_monto = 10 if currency == "MXN" else 15
            monto_val = row[col_monto] if len(row) > col_monto else None
            fecha_val = row[16] if len(row) > 16 else None
            bank_val = row[18] if len(row) > 18 else None
            edo_cta_val = row[20] if len(row) > 20 else None

        rfc_clean = str(rfc_val).strip()
        if moneda_val == currency or rfc_clean in rfc_excluidos or fecha_val is None:
            continue

        dia_a_buscar = formatear_fecha(fecha_val, type_bank)

        if monto_val is not None and bank_val == type_bank:
            try:
                # Convertimos a float y luego a int para quitar decimales del "buscado"
                monto_buscable = str(int(float(monto_val)))
                monto_original = f"${float(monto_val):,.2f}"
            except (ValueError, TypeError):
                continue

            datos_monto.append(
                {
                    "dia_a_buscar": dia_a_buscar,
                    "monto_limpio": monto_buscable,
                    "monto_original": monto_original,
                    "referencia": edo_cta_val,
                }
            )
    wb_iva.close()
    return datos_monto


def contiene_flag_break(texto_de_la_pagina, flag_break):
    texto_unido = " ".join(texto_de_la_pagina)
    return any(flag in texto_unido for flag in flag_break)


def resaltar_por_fecha_y_monto(pdf_entrada, pdf_salida, items_a_buscar):
    print("🤖 Inicializando motor OCR...")
    # reader = easyocr.Reader(["es", "en"])
    reader = crear_reader()
    doc = fitz.open(pdf_entrada)
    encontrados_totales = 0

    print(f"🔍 Iniciando búsqueda de {len(items_a_buscar)} transacciones...")

    for num_pagina in range(len(doc)):
        pagina = doc[num_pagina]
        print(f"📄 Procesando página {num_pagina + 1}...", end="\r")

        pix = pagina.get_pixmap(dpi=300)
        img_data = np.frombuffer(pix.samples, dtype=np.uint8).reshape(
            (pix.h, pix.w, pix.n)
        )
        img_rgb = (
            cv2.cvtColor(img_data, cv2.COLOR_RGBA2RGB)
            if pix.n == 4
            else cv2.cvtColor(img_data, cv2.COLOR_RGB2BGR)
        )

        resultados = reader.readtext(img_rgb)
        texto_completo_pag = [res[1].upper() for res in resultados]

        # 1. Banderas de salida
        flag_break = [
            "INFORMACIÓN SPEI",
            "SE INFORMA QUE LA COMISIÓN",
            "S ENVIADOS DURANTE EL PERIODO",
        ]
        if contiene_flag_break(texto_completo_pag, flag_break):
            print("\n🛑 Flag de salida detectada.")
            break

        # 2. Búsqueda
        for bbox, texto_ocr, confidence in resultados:
            # Limpiamos el texto del OCR para tener solo números (ej: "\$1,815.00" -> "181500")
            texto_ocr_limpio = "".join(c for c in texto_ocr if c.isdigit())

            if len(texto_ocr_limpio) < 1:
                continue

            for item in items_a_buscar:
                monto_buscado = item["monto_limpio"]  # Ej: "990"
                print(monto_buscado)

                # --- LÓGICA DE VALIDACIÓN ESTRICTA ---
                if texto_ocr_limpio.startswith(monto_buscado):
                    # Calculamos el sobrante (posibles centavos)
                    sobrante = texto_ocr_limpio[len(monto_buscado) :]

                    # Solo es válido si sobran 0, 1 o 2 dígitos (centavos)
                    # Si sobran 3 o más, es otro número (ej: buscaba 181 y encontró 181500)
                    if len(sobrante) <= 2:
                        # Validamos que el día de la transacción esté en la misma página
                        if any(
                            item["dia_a_buscar"] == t.zfill(2)
                            for t in texto_completo_pag
                            if t.isdigit()
                        ):
                            encontrados_totales += 1
                            print(
                                f"\n🎯 Match: Día {item['dia_a_buscar']} - Monto {item['monto_original']} (OCR: {texto_ocr})"
                            )

                            # Dibujar resaltado
                            y_min = min([p[1] for p in bbox]) / (
                                pix.h / pagina.rect.height
                            )
                            y_max = max([p[1] for p in bbox]) / (
                                pix.h / pagina.rect.height
                            )
                            rect_renglon = fitz.Rect(
                                15, y_min - 1, pagina.rect.width - 15, y_max + 1
                            )

                            resaltado = pagina.add_highlight_annot(rect_renglon)
                            resaltado.set_colors({"stroke": [1, 1, 0]})
                            resaltado.update()

                            if item["referencia"]:
                                pos_x = pagina.rect.width - 130
                                pagina.insert_text(
                                    fitz.Point(pos_x, y_max - 2),
                                    str(item["referencia"]),
                                    fontsize=7,
                                    color=(1, 0, 0),
                                    fontname="helv",
                                )
                            break

    if encontrados_totales > 0:
        doc.save(pdf_salida, garbage=4, deflate=True)
        print(
            f"\n✅ Proceso terminado. Se resaltaron {encontrados_totales} transacciones."
        )
    else:
        print("\n❌ No se encontraron coincidencias.")
    doc.close()


def procesar_por_suff(suff, path, file_name, config):
    path = Path(path).resolve()  # convierte a Path y resuelve rutas absolutas

    suff_key = suff.lower()
    if suff_key not in config:
        return

    datos = config[suff_key]
    banks = datos.get("banks", [])
    name_sheet = datos.get("sheet_name", "")

    archivo_base = path / "{file_name}.xlsx"

    if not archivo_base.exists() or not egresos_path.exists():
        faltante = "diot.xlsx" if not diot_path.exists() else "egresos.xlsx"
        msg = f"No se encontró el archivo necesario: {faltante}"
        self.text_console_log(msg, "ERROR")
        # Este RAISE detiene la función AQUÍ MISMO y va directo al except de la GUI
        raise FileNotFoundError(msg)

    for bank in banks:
        bank_file = bank["file"]
        archivo_input = f"{path}/{bank_file}.pdf"
        archivo_output = f"{path}/{bank_file}_Underlined.pdf"
        type_bank = extraer_caracteres_alfabeticos(bank_file)

        transacciones = obtener_monto_por_poliza(
            archivo_base=archivo_base,
            currency=bank["currency"],
            version=suff,
            nombre_hoja=name_sheet,
            type_bank=type_bank,
        )

        print(transacciones)

        resaltar_por_fecha_y_monto(archivo_input, archivo_output, transacciones)


# ====================== CONFIGURACIÓN ======================
# suff = "v16"  # ← Cambia aquí: "iva", "v0" o "v16"
# mounth = "sep"
# year = "2025"  # ← Nombre del PDF sin extensión (ajústalo)

configuracion = {
    "iva": {
        "sheet_name": "",
        "banks": [
            {"file": "BASE2018", "currency": "MXN"},
            {"file": "HSBC8881", "currency": "USD"},
        ],
    },
    "v0": {
        "sheet_name": "",
        "banks": [
            {"file": "HSBC5430", "currency": "USD"},
        ],
    },
    "v16": {
        "sheet_name": "",
        "banks": [
            {"file": "HSBC5430", "currency": "USD"},
            {"file": "BASE2018", "currency": "USD"},
            {"file": "HSBC8881", "currency": "MXN"},
        ],
    },
}


def main_underline_bank_process(suff, path, file_name):
    procesar_por_suff(suff, path, file_name, configuracion)
